#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Japanese PowerPoint Generation Example

Demonstrates generating a PowerPoint presentation with Japanese text content
using IBM Plex Sans JP font while maintaining IBM Carbon Design System compliance.

This example:
1. Selects slides from the IBM Consulting Offerings template
2. Replaces English text with Japanese content
3. Sets IBM Plex Sans JP font for Japanese text
4. Validates encoding and font assignments
5. Generates a complete Japanese presentation

Usage:
    python japanese_example.py
"""

import shutil
import zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation

# Import utilities
import sys
sys.path.insert(0, str(Path(__file__).parent))
from japanese_utils import contains_japanese, analyze_japanese_content, print_analysis_report
from font_utils import set_japanese_fonts

# Paths
SCRIPT_DIR = Path(__file__).parent
SKILL_DIR = SCRIPT_DIR.parent
TEMPLATE_PATH = SKILL_DIR / "template" / "Example_layouts_IBM_Consulting_Offerings_template_v_1_0_011025.pptx"
OUTPUT_DIR = Path("slides/japanese-example")
OUTPUT_PATH = OUTPUT_DIR / "IBM_Japanese_Presentation.pptx"

# XML namespaces
NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}

# Selected slides from template
SELECTED_SLIDES = [2, 5, 6]  # Title slide, Image+Title slide, Quote slide

# Japanese content replacements (mapped to visible slide order after filtering)
SLIDE_REPLACEMENTS = {
    1: {  # Visible slide 1 (template slide 2 - Title slide)
        "L1 (Halo)": "IBMコンサルティング",
    },
    2: {  # Visible slide 2 (template slide 5 - Image+Title slide)
        "Developing an AI Strategy": "AI戦略の開発",
        "Project Summary 2024": "プロジェクト概要 2024年",
        "Jack Kirby": "田中太郎",
    },
    3: {  # Visible slide 3 (template slide 6 - Quote slide)
        "The number 1 request from deal teams is more client stories.": "営業チームからの最も多い要望は、より多くのクライアント事例です。",
    }
}

# Speaker notes in Japanese
SPEAKER_NOTES = {
    1: "このプレゼンテーションでは、IBMのデジタルトランスフォーメーション戦略について説明します。",
    2: "プロジェクトの主な目的と期間、予算について説明してください。クライアントのニーズに焦点を当てます。",
    3: "次のステップとして、要件定義を完了し、システム設計を開始します。開発チームの編成も並行して進めます。",
}


def delete_slide(prs: Presentation, slide_index: int) -> None:
    """
    Delete slide by zero-based index using python-pptx package API.
    
    Args:
        prs: Presentation object
        slide_index: Zero-based slide index to delete
    """
    slide_id_list = prs.slides._sldIdLst
    slide_id = slide_id_list[slide_index]
    rel_id = slide_id.rId
    prs.part.drop_rel(rel_id)
    slide_id_list.remove(slide_id)


def keep_only_slides(prs: Presentation, selected_slide_numbers: list[int]) -> None:
    """
    Keep only specified 1-based slide numbers from source deck.
    
    Args:
        prs: Presentation object
        selected_slide_numbers: List of 1-based slide numbers to keep
    """
    selected_indexes = {slide_number - 1 for slide_number in selected_slide_numbers}
    for index in reversed(range(len(prs.slides))):
        if index not in selected_indexes:
            delete_slide(prs, index)


def replace_in_slide_xml(root: etree._Element, replacements: dict) -> dict:
    """
    Replace text in slide XML at paragraph level.
    
    Args:
        root: XML root element
        replacements: Dictionary mapping old text to new text
        
    Returns:
        Dictionary mapping old text to replacement count
    """
    counts = {old: 0 for old in replacements}
    
    for paragraph in root.xpath('.//a:p', namespaces=NS):
        nodes = paragraph.xpath('.//a:t', namespaces=NS)
        if not nodes:
            continue
        
        # Join all text nodes in paragraph
        original = ''.join(node.text or '' for node in nodes)
        updated = original
        
        # Apply all replacements
        for old, new in replacements.items():
            if old in updated:
                updated = updated.replace(old, new)
                counts[old] += 1
        
        # Update text if changed
        if updated != original:
            nodes[0].text = updated
            for node in nodes[1:]:
                node.text = ''
    
    return counts


def replace_text_in_presentation(pptx_path: Path, replacements_by_slide: dict) -> dict:
    """
    Replace text in presentation using direct XML manipulation.
    
    Args:
        pptx_path: Path to PowerPoint file
        replacements_by_slide: Dictionary mapping visible slide numbers to replacement dictionaries
        
    Returns:
        Dictionary with replacement statistics
    """
    modified_parts = {}
    all_counts = {}
    
    with zipfile.ZipFile(pptx_path, 'r') as zin:
        # Get slide parts in order
        slide_parts = sorted([n for n in zin.namelist() if 'ppt/slides/slide' in n and n.endswith('.xml')])
        
        for visible_slide_num, replacements in replacements_by_slide.items():
            if visible_slide_num > len(slide_parts):
                print(f"Warning: Slide {visible_slide_num} does not exist (only {len(slide_parts)} slides)")
                continue
            
            slide_part = slide_parts[visible_slide_num - 1]
            root = etree.fromstring(zin.read(slide_part))
            
            # Replace text
            counts = replace_in_slide_xml(root, replacements)
            all_counts[visible_slide_num] = counts
            
            # Check for unmatched replacements
            missing = [old for old, count in counts.items() if count == 0]
            if missing:
                print(f"Warning: Slide {visible_slide_num} - Unmatched replacements: {missing}")
            
            # Store modified XML
            modified_parts[slide_part] = etree.tostring(
                root,
                encoding='UTF-8',
                xml_declaration=True
            )
    
    # Write modified presentation
    tmp_path = pptx_path.with_suffix('.tmp.pptx')
    with zipfile.ZipFile(pptx_path, 'r') as zin:
        with zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = modified_parts.get(item.filename, zin.read(item.filename))
                zout.writestr(item, data)
    
    tmp_path.replace(pptx_path)
    
    return all_counts


def add_speaker_notes(pptx_path: Path, notes_by_slide: dict) -> None:
    """
    Add speaker notes to slides using python-pptx.
    
    Args:
        pptx_path: Path to PowerPoint file
        notes_by_slide: Dictionary mapping slide numbers to notes text
    """
    prs = Presentation(str(pptx_path))
    
    for slide_num, notes_text in notes_by_slide.items():
        if slide_num <= len(prs.slides):
            slide = prs.slides[slide_num - 1]
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = notes_text
    
    prs.save(str(pptx_path))


def validate_presentation(pptx_path: Path, expected_slides: int) -> bool:
    """
    Validate generated presentation.
    
    Args:
        pptx_path: Path to PowerPoint file
        expected_slides: Expected number of slides
        
    Returns:
        True if validation passes
    """
    print("\n" + "="*60)
    print("Validation")
    print("="*60)
    
    # 1. File exists
    if not pptx_path.exists():
        print(f"✗ File not found: {pptx_path}")
        return False
    print(f"✓ File exists: {pptx_path}")
    
    # 2. ZIP integrity
    try:
        with zipfile.ZipFile(pptx_path, 'r') as z:
            if z.testzip() is not None:
                print("✗ Corrupted ZIP package")
                return False
        print("✓ ZIP integrity OK")
    except Exception as e:
        print(f"✗ ZIP validation failed: {e}")
        return False
    
    # 3. Slide count
    try:
        prs = Presentation(str(pptx_path))
        actual_count = len(prs.slides)
        if actual_count != expected_slides:
            print(f"✗ Slide count mismatch: expected {expected_slides}, got {actual_count}")
            return False
        print(f"✓ Slide count correct: {actual_count}")
    except Exception as e:
        print(f"✗ Failed to load presentation: {e}")
        return False
    
    # 4. Japanese content validation
    try:
        analysis = analyze_japanese_content(pptx_path)
        if analysis['slides_with_japanese'] == 0:
            print("✗ No Japanese content found")
            return False
        print(f"✓ Japanese content found in {analysis['slides_with_japanese']} slides")
        
        if not analysis['encoding_valid']:
            print("✗ UTF-8 encoding validation failed")
            return False
        print("✓ UTF-8 encoding valid")
        
        if not analysis['fonts_valid']:
            print("✗ Font validation failed (IBM Plex Sans JP not set)")
            return False
        print("✓ IBM Plex Sans JP font correctly assigned")
        
    except Exception as e:
        print(f"✗ Content validation failed: {e}")
        return False
    
    print("\n✓ All validation checks passed")
    print("="*60 + "\n")
    
    return True


def main():
    """Generate Japanese presentation example."""
    print("\n" + "="*60)
    print("Japanese PowerPoint Generation Example")
    print("="*60)
    print(f"\nTemplate: {TEMPLATE_PATH.name}")
    print(f"Output: {OUTPUT_PATH}")
    print(f"Selected slides: {SELECTED_SLIDES}")
    
    # Verify template exists
    if not TEMPLATE_PATH.exists():
        print(f"\n✗ Error: Template not found: {TEMPLATE_PATH}")
        return 1
    
    # Create output directory
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"\n✓ Created output directory: {OUTPUT_DIR}")
    
    # Step 1: Load template and filter slides
    print("\n" + "-"*60)
    print("Step 1: Filter slides")
    print("-"*60)
    prs = Presentation(str(TEMPLATE_PATH))
    original_count = len(prs.slides)
    print(f"Original slide count: {original_count}")
    
    keep_only_slides(prs, SELECTED_SLIDES)
    print(f"Filtered to {len(prs.slides)} slides: {SELECTED_SLIDES}")
    
    # Save filtered presentation
    prs.save(str(OUTPUT_PATH))
    print(f"✓ Saved filtered presentation")
    
    # Step 2: Replace text with Japanese content
    print("\n" + "-"*60)
    print("Step 2: Replace text with Japanese content")
    print("-"*60)
    replacement_counts = replace_text_in_presentation(OUTPUT_PATH, SLIDE_REPLACEMENTS)
    
    for slide_num, counts in replacement_counts.items():
        total = sum(counts.values())
        print(f"Slide {slide_num}: {total} replacements")
        for old, count in counts.items():
            if count > 0:
                preview = old[:30] + "..." if len(old) > 30 else old
                print(f"  - '{preview}' → {count}x")
    
    print("✓ Text replacement complete")
    
    # Step 3: Set IBM Plex Sans JP font
    print("\n" + "-"*60)
    print("Step 3: Set IBM Plex Sans JP font")
    print("-"*60)
    font_results = set_japanese_fonts(OUTPUT_PATH)
    print(f"Processed {font_results['slides_processed']} slides")
    print(f"Updated {font_results['runs_updated']} text runs in {font_results['slides_with_updates']} slides")
    print("✓ Font assignment complete")
    
    # Step 4: Add speaker notes
    print("\n" + "-"*60)
    print("Step 4: Add speaker notes")
    print("-"*60)
    add_speaker_notes(OUTPUT_PATH, SPEAKER_NOTES)
    print(f"Added notes to {len(SPEAKER_NOTES)} slides")
    print("✓ Speaker notes added")
    
    # Step 5: Validate
    print("\n" + "-"*60)
    print("Step 5: Validate presentation")
    print("-"*60)
    if not validate_presentation(OUTPUT_PATH, len(SELECTED_SLIDES)):
        print("\n✗ Validation failed")
        return 1
    
    # Step 6: Analysis report
    print("\n" + "-"*60)
    print("Step 6: Content analysis")
    print("-"*60)
    analysis = analyze_japanese_content(OUTPUT_PATH)
    print_analysis_report(analysis)
    
    # Summary
    print("="*60)
    print("Generation Complete")
    print("="*60)
    print(f"\n✓ Generated: {OUTPUT_PATH}")
    print(f"✓ Slides: {len(SELECTED_SLIDES)}")
    print(f"✓ Japanese text: {analysis['slides_with_japanese']} slides")
    print(f"✓ Font: IBM Plex Sans JP")
    print(f"✓ Encoding: UTF-8")
    print(f"✓ Speaker notes: {len(SPEAKER_NOTES)} slides")
    
    print(f"\nOpen the presentation:")
    print(f"  open {OUTPUT_PATH}")
    print()
    
    return 0


if __name__ == "__main__":
    sys.exit(main())
