#!/usr/bin/env python3
"""Sample direct Open XML deck generator for the PPT template slide workflow.

This file is intentionally a compact, auditable example rather than a shared
replacement library. Copy this pattern into a deck-specific generator and keep
the selected slide parts, replacement map, explicit clears, and validation in
that generator so each deck is easy to inspect and debug.

CRITICAL: This script demonstrates the correct workflow for template immutability:
1. Load template (read-only)
2. Modify in memory
3. Save to NEW location (never overwrite template)

Example:
    uv run python .bob/skills/ppt-template-slide-workflow/scripts/direct_xml_replace_sample.py
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

# Add scripts directory to path for template protection
sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "scripts"))
from template_protection import validate_write_path, safe_save_presentation


ROOT = Path(__file__).resolve().parents[4]
TEMPLATE_PATH = (
    ROOT
    / ".agents"
    / "skills"
    / "ppt-template-slide-workflow"
    / "template"
    / "Example_layouts_IBM_Consulting_Offerings_template_v_1_0_011025.pptx"
)
WORKSPACE_DIR = ROOT / "slides" / "direct-xml-sample"
OUTPUT_PATH = WORKSPACE_DIR / "direct_xml_sample_deck.pptx"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

NS = {"p": P_NS, "a": A_NS, "r": R_NS, "rel": REL_NS}

SELECTED_SLIDES = [3, 5]

# Replacement maps are keyed by visible slide number after slide filtering.
SLIDE_REPLACEMENTS = {
    1: {
        "Consulting Offerings Presentation Template": "Direct XML Replacement Sample",
        "Firstname Lastname": "Template Workflow Skill",
        "Job Title": "Open XML generation pattern",
        "email.address@ibm.com": "",
        "123.456.7890": "",
    },
    2: {
        "Developing an AI Strategy": "Why direct XML replacement?",
        "Project Summary 2024": "Edit ppt/slides/slideN.xml directly\nAssert every required replacement\nPreserve template geometry and styling\nValidate package integrity before delivery",
        "Jack Kirby": "",
    },
}


def delete_slide(prs: Presentation, slide_index: int) -> None:
    """Delete a slide by zero-based index using python-pptx package APIs."""

    slide_id_list = prs.slides._sldIdLst
    slide_id = slide_id_list[slide_index]
    relationship_id = slide_id.rId
    prs.part.drop_rel(relationship_id)
    slide_id_list.remove(slide_id)


def keep_only_slides(prs: Presentation, selected_slide_numbers: list[int]) -> None:
    """Keep only 1-based slide numbers from the source template."""

    selected_indexes = {slide_number - 1 for slide_number in selected_slide_numbers}
    for index in reversed(range(len(prs.slides))):
        if index not in selected_indexes:
            delete_slide(prs, index)


def ordered_visible_slide_parts(pptx_path: Path) -> list[str]:
    """Return visible slide XML parts in presentation order."""

    with zipfile.ZipFile(pptx_path) as deck:
        presentation = etree.fromstring(deck.read("ppt/presentation.xml"))
        relationships = etree.fromstring(deck.read("ppt/_rels/presentation.xml.rels"))
        targets = {
            relationship.get("Id"): relationship.get("Target")
            for relationship in relationships.xpath("./rel:Relationship", namespaces=NS)
        }

        slide_parts: list[str] = []
        for slide_id in presentation.xpath(".//p:sldIdLst/p:sldId", namespaces=NS):
            relationship_id = slide_id.get(f"{{{R_NS}}}id")
            target = targets[relationship_id]
            if target.startswith("../"):
                target = target[3:]
            if not target.startswith("ppt/"):
                target = f"ppt/{target}"
            slide_parts.append(target)
        return slide_parts


def paragraph_text_nodes(root: etree._Element) -> list[list[etree._Element]]:
    """Group DrawingML text nodes by paragraph."""

    groups: list[list[etree._Element]] = []
    for paragraph in root.xpath(".//a:p", namespaces=NS):
        nodes = paragraph.xpath(".//a:t", namespaces=NS)
        if nodes:
            groups.append(nodes)
    return groups


def set_paragraph_text(nodes: list[etree._Element], text: str) -> None:
    """Set paragraph text while preserving the first run's formatting."""

    nodes[0].text = text
    for node in nodes[1:]:
        node.text = ""


def replace_in_slide_xml(root: etree._Element, replacements: dict[str, str]) -> dict[str, int]:
    """Replace paragraph text directly in one slide XML tree."""

    counts = {old_text: 0 for old_text in replacements}
    for nodes in paragraph_text_nodes(root):
        original = "".join(node.text or "" for node in nodes)
        updated = original
        for old_text, new_text in replacements.items():
            if old_text in updated:
                updated = updated.replace(old_text, new_text)
                counts[old_text] += 1
        if updated != original:
            set_paragraph_text(nodes, updated)
    return counts


def rewrite_modified_parts(pptx_path: Path, modified_parts: dict[str, bytes]) -> None:
    """Rewrite a PPTX package with selected XML parts replaced."""

    temporary_path = pptx_path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(pptx_path, "r") as source_deck:
        with zipfile.ZipFile(temporary_path, "w") as output_deck:
            written: set[str] = set()
            for item in source_deck.infolist():
                if item.filename in written:
                    continue
                data = modified_parts.get(item.filename, source_deck.read(item.filename))
                output_deck.writestr(item, data)
                written.add(item.filename)
    temporary_path.replace(pptx_path)


def apply_direct_xml_replacements(pptx_path: Path) -> None:
    """Apply direct XML replacements and fail on unmatched required keys."""

    modified_parts: dict[str, bytes] = {}
    visible_parts = ordered_visible_slide_parts(pptx_path)

    with zipfile.ZipFile(pptx_path, "r") as deck:
        for visible_slide, slide_part in enumerate(visible_parts, 1):
            replacements = SLIDE_REPLACEMENTS[visible_slide]
            root = etree.fromstring(deck.read(slide_part))
            counts = replace_in_slide_xml(root, replacements)
            missing = [old_text for old_text, count in counts.items() if count == 0]
            if missing:
                raise ValueError(f"Unmatched replacements on visible slide {visible_slide}: {missing}")
            modified_parts[slide_part] = etree.tostring(
                root,
                encoding="UTF-8",
                xml_declaration=True,
            )

    rewrite_modified_parts(pptx_path, modified_parts)


def validate_package(pptx_path: Path) -> None:
    """Validate ZIP integrity and duplicate entries."""

    with zipfile.ZipFile(pptx_path) as deck:
        corrupt_entry = deck.testzip()
        if corrupt_entry is not None:
            raise ValueError(f"Corrupt PPTX entry: {corrupt_entry}")
        duplicate_entries = [name for name in deck.namelist() if deck.namelist().count(name) > 1]
        if duplicate_entries:
            raise ValueError(f"Duplicate PPTX entries: {sorted(set(duplicate_entries))}")


def main() -> None:
    """Generate a tiny sample deck using direct XML replacement.
    
    CRITICAL: Demonstrates template immutability principle:
    - Load template (read-only)
    - Modify in memory
    - Save to NEW location with validation
    """

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    WORKSPACE_DIR.mkdir(parents=True, exist_ok=True)

    # ✅ CORRECT: Load template (read-only)
    presentation = Presentation(str(TEMPLATE_PATH))
    
    # ✅ CORRECT: Modify in memory
    keep_only_slides(presentation, SELECTED_SLIDES)
    
    # ✅ CORRECT: Validate output path before saving
    validate_write_path(OUTPUT_PATH)
    
    # ✅ CORRECT: Save to NEW location (never overwrite template)
    presentation.save(str(OUTPUT_PATH))

    apply_direct_xml_replacements(OUTPUT_PATH)
    validate_package(OUTPUT_PATH)

    print(f"Generated sample deck: {OUTPUT_PATH}")
    print(f"Selected source slides: {SELECTED_SLIDES}")
    print("Replacement method: direct ppt/slides/slideN.xml editing")


if __name__ == "__main__":
    main()
